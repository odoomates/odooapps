"""Parser of the CSV, TSV and Excel files, with the map of the columns of the bank.

The layout of such a file changes from one bank to the next: it is given by ``options['map']``, the dict of
``account.statement.import.map._get_map_options()``. Without a map, the file cannot be read and the parser
gives it back to the next one.
"""
import csv
import datetime
import io
import re

try:
    import openpyxl
except ImportError:
    openpyxl = None

# the extensions we read
EXCEL_EXTENSIONS = ('.xlsx', '.xlsm', '.xltx', '.xltm', '.xls')
CSV_EXTENSIONS = ('.csv', '.tsv', '.tab', '.txt')
# the extensions of the other parsers, never ours
FOREIGN_EXTENSIONS = ('.xml', '.ofx', '.qfx', '.qif', '.sta', '.940', '.mt940', '.json', '.pdf', '.zip')

ZIP_SIGNATURE = b'PK\x03\x04'
OLE_SIGNATURE = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'
# the first characters of the files of the other formats: XML/CAMT, OFX, QIF, MT940, PDF
FOREIGN_HEAD = re.compile(
    r'^(<|%PDF|!Type:|!Account|OFXHEADER|\{1:|:20:|:940:|-{0,3}\s*:\d{2}[A-Z]?:)', re.IGNORECASE)


def parse(data, filename, options):
    """ Read a CSV, TSV or Excel file with the map of the file. See ``parsers/__init__.py``. """
    file_map = (options or {}).get('map')
    if not file_map:
        # without a map we cannot know the layout of the file: let the other parsers try
        return None
    if not data:
        return None

    name = (filename or '').lower()
    extension = name[name.rfind('.'):] if '.' in name else ''
    looks_like_excel = data[:4] == ZIP_SIGNATURE or data[:8] == OLE_SIGNATURE

    if extension in FOREIGN_EXTENSIONS and not looks_like_excel:
        return None
    if not looks_like_excel and _is_foreign_content(data):
        return None

    date_format = (options or {}).get('date_format') or file_map.get('date_format')
    if not date_format:
        raise ValueError('The map %r has no date format: the dates of the file cannot be read.'
                         % file_map.get('name'))

    if looks_like_excel or extension in EXCEL_EXTENSIONS:
        rows = _read_excel(data, file_map, extension)
        if rows is None:
            return None
    else:
        rows = _read_csv(data, file_map, extension)
        if rows is None:
            return None
    return _build_statements(rows, file_map, date_format)


# -------------------------------------------------------------------------
# Recognition of the format
# -------------------------------------------------------------------------
def _is_foreign_content(data):
    """ :return: True when the bytes are clearly of another format (XML/CAMT, OFX, QIF, MT940, PDF) """
    head = data[:2048].decode('utf-8', errors='ignore').lstrip('﻿ \t\r\n')
    if not head.strip():
        return False
    if FOREIGN_HEAD.match(head.strip()):
        return True
    upper = head.upper()
    return '<OFX' in upper or 'OFXHEADER' in upper


# -------------------------------------------------------------------------
# Reading of the rows
# -------------------------------------------------------------------------
def _read_csv(data, file_map, extension):
    """ :return: the rows of the CSV file as lists of strings, or None when it is not text """
    encoding = file_map.get('encoding') or 'utf-8'
    try:
        text = data.decode(encoding)
    except (UnicodeDecodeError, LookupError) as error:
        if extension in CSV_EXTENSIONS:
            raise ValueError('The file cannot be read with the encoding %s of the map %r: %s'
                             % (encoding, file_map.get('name'), error))
        return None  # binary, of a format we do not know: the next parser may
    if text.startswith('﻿'):
        text = text[1:]

    delimiter = _delimiter(file_map, extension)
    quotechar = (file_map.get('quotechar') or '"')[:1] or '"'
    try:
        return list(csv.reader(io.StringIO(text, newline=''), delimiter=delimiter, quotechar=quotechar))
    except csv.Error as error:
        raise ValueError('The file is not a valid CSV file: %s' % error)


def _delimiter(file_map, extension):
    delimiter = file_map.get('delimiter') or ','
    if delimiter.strip().lower() in ('tab', '\\t') or delimiter == '\t':
        return '\t'
    if extension in ('.tsv', '.tab') and delimiter == ',':
        return '\t'  # the map keeps the default comma but the file is a tabulated one
    return delimiter[:1]


def _read_excel(data, file_map, extension):
    """ :return: the rows of the sheet of the map as lists of values, or None when it is not a workbook """
    if openpyxl is None:
        raise ValueError('Reading an Excel file needs the Python library openpyxl. Install it, or export '
                         'the statement of your bank in CSV.')
    if data[:8] == OLE_SIGNATURE:
        raise ValueError('The old Excel format (.xls) cannot be read. Open the file and save it as .xlsx, '
                         'or export the statement of your bank in CSV.')
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as error:
        if extension in EXCEL_EXTENSIONS:
            raise ValueError('The Excel file cannot be opened: %s' % error)
        return None  # a zip file of something else: the next parser may
    try:
        sheet = _pick_sheet(workbook, file_map)
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _pick_sheet(workbook, file_map):
    sheet = (file_map.get('sheet') or '').strip()
    if not sheet:
        return workbook.worksheets[0]
    if sheet in workbook.sheetnames:
        return workbook[sheet]
    if sheet.isdigit():
        number = int(sheet)
        if 1 <= number <= len(workbook.worksheets):
            return workbook.worksheets[number - 1]
    raise ValueError('The file has no sheet %r. Its sheets are: %s.' % (sheet, ', '.join(workbook.sheetnames)))


# -------------------------------------------------------------------------
# Reading of the transactions
# -------------------------------------------------------------------------
def _build_statements(rows, file_map, date_format):
    columns = file_map.get('columns') or {}
    has_header = file_map.get('has_header', True)
    skip_rows = int(file_map.get('skip_rows') or 0)

    # the number of the first row of ``rows`` in the file, to name the rows in the messages
    first_row_number = skip_rows + 1
    body = rows[skip_rows:]
    header = None
    if has_header:
        while body and _is_empty_row(body[0]):
            body = body[1:]
            first_row_number += 1
        if not body:
            raise ValueError('The file has no header row: it is empty, or too many rows are skipped.')
        header = [_as_text(cell) for cell in body[0]]
        body = body[1:]
        first_row_number += 1

    indexes = {key: _column_index(key, spec, header, has_header, file_map)
               for key, spec in columns.items()}
    _check_columns(indexes, header, has_header)

    transactions = []
    currency_code = None
    for offset, row in enumerate(body):
        row_number = first_row_number + offset
        if _is_empty_row(row):
            continue
        transaction, row_currency = _build_transaction(row, indexes, file_map, date_format, row_number)
        transactions.append(transaction)
        currency_code = currency_code or row_currency

    if not transactions:
        raise ValueError('The file holds no transaction. Check the rows to skip and the header of the map %r.'
                         % file_map.get('name'))

    dates = [transaction['date'] for transaction in transactions]
    return [{
        'name': None,
        'date': max(dates),
        'balance_start': None,
        'balance_end_real': None,
        'currency_code': currency_code,
        'account_number': None,
        'transactions': transactions,
    }]


def _check_columns(indexes, header, has_header):
    if indexes.get('date') is None:
        raise ValueError('The date column of the map is missing from the file.%s'
                         % _header_hint(header, has_header))
    if indexes.get('amount') is None and (indexes.get('debit') is None or indexes.get('credit') is None):
        raise ValueError('The amount column, or the debit and the credit columns, of the map are missing '
                         'from the file.%s' % _header_hint(header, has_header))


def _header_hint(header, has_header):
    if has_header and header:
        return ' The columns of the file are: %s.' % ', '.join('%r' % name for name in header if name)
    return ''


def _column_index(key, spec, header, has_header, file_map):
    """ :return: the 0-based index of a column of the map, by its name or by its number, None when unset """
    spec = (spec or '').strip()
    if not spec:
        return None
    if has_header and header:
        wanted = spec.lower()
        for index, name in enumerate(header):
            if name.strip().lower() == wanted:
                return index
    if spec.isdigit() and int(spec) > 0:
        return int(spec) - 1
    if has_header:
        return None
    raise ValueError('The file has no header, so the column %s of the map %r must be the number of the '
                     'column, not %r.' % (key, file_map.get('name'), spec))


def _build_transaction(row, indexes, file_map, date_format, row_number):
    date = _to_date(_cell(row, indexes.get('date')), date_format, row_number)
    amount = _to_amount(row, indexes, file_map, row_number)
    label = _as_text(_cell(row, indexes.get('label')))
    ref = _as_text(_cell(row, indexes.get('ref')))
    unique_id = _as_text(_cell(row, indexes.get('unique_id')))
    currency_code = _as_text(_cell(row, indexes.get('currency'))).upper() or None
    return {
        'date': date,
        'payment_ref': label or ref or '/',
        'amount': amount,
        'unique_import_id': unique_id or None,
        'partner_name': _as_text(_cell(row, indexes.get('partner'))) or None,
        'account_number': _as_text(_cell(row, indexes.get('account_number'))) or None,
        'ref': ref or None,
        'narration': _as_text(_cell(row, indexes.get('narration'))) or None,
        'transaction_type': None,
    }, currency_code


def _to_amount(row, indexes, file_map, row_number):
    if indexes.get('amount') is not None:
        return _to_float(_cell(row, indexes['amount']), file_map, row_number, 'amount')
    debit = _to_float(_cell(row, indexes.get('debit')), file_map, row_number, 'debit')
    credit = _to_float(_cell(row, indexes.get('credit')), file_map, row_number, 'credit')
    if file_map.get('debit_positive', True):
        # the bank writes the money leaving the account as a positive number
        debit = -abs(debit)
    return credit + debit


# -------------------------------------------------------------------------
# Values
# -------------------------------------------------------------------------
def _cell(row, index):
    if index is None or index < 0 or index >= len(row):
        return None
    return row[index]


def _as_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel gives 1234.0 for a number written 1234
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    return str(value).strip()


def _is_empty_row(row):
    return all(_as_text(cell) == '' for cell in row)


def _to_date(value, date_format, row_number):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = _as_text(value)
    if not text:
        raise ValueError('Row %s of the file has no date.' % row_number)
    try:
        return datetime.datetime.strptime(text, date_format).date()
    except ValueError:
        raise ValueError('Row %s of the file: the date %r does not match the format %r of the map.'
                         % (row_number, text, date_format))


def _to_float(value, file_map, row_number, column):
    if value is None or value == '':
        return 0.0
    if isinstance(value, bool):
        raise ValueError('Row %s of the file: the %s %r is not a number.' % (row_number, column, value))
    if isinstance(value, (int, float)):
        return float(value)

    text = _as_text(value)
    if not text:
        return 0.0
    negative = False
    if text.startswith('(') and text.endswith(')'):  # (1 234,56) is a negative amount
        negative, text = True, text[1:-1].strip()
    if text.endswith('-'):  # 1 234,56- is a negative amount too
        negative, text = True, text[:-1].strip()

    thousands = file_map.get('thousands_separator') or ''
    decimal = file_map.get('decimal_separator') or '.'
    cleaned = text.replace(' ', '').replace(' ', '')
    if thousands:
        cleaned = cleaned.replace(thousands, '')
    if decimal != '.':
        cleaned = cleaned.replace('.', '') if not thousands else cleaned
        cleaned = cleaned.replace(decimal, '.')
    cleaned = cleaned.replace('+', '')
    try:
        amount = float(cleaned)
    except ValueError:
        raise ValueError('Row %s of the file: the %s %r is not a number. Check the decimal separator and '
                         'the thousands separator of the map.' % (row_number, column, text))
    return -amount if negative else amount
